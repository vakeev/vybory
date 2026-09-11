#!/usr/bin/env python3
"""Generate site pages from candidates.xlsx and voprosy.txt.

Pages:
  index.html                  root page: «Все вопросы» / «Все кандидаты» nav + list of regions
  voprosy-kandidatam.html     question list with copy buttons
  voprosy.js                  question bank as JS data (window.VOPROSY)
  regions/<slug>.html         candidates of one region, each with a random-question button

Region slug -> name mapping is read from the existing region pages.
"""
import glob
import html
import json
import os
import re

from openpyxl import load_workbook

REPO = "candidates.xlsx"
REPO_SHEET = "Кандидаты по округам"
QUESTIONS_FILE = "voprosy.txt"
CONTACT_COLS = ["ВКонтакте", "Telegram", "запрещённая сеть", "Одноклассники", "MAX", "Сайт", "Дзен"]
ALLOWED_PARTIES = {"Единая Россия", "КПРФ", "ЛДПР", "Справедливая Россия", "Новые люди"}

HEAD = "<!DOCTYPE html>\n<html lang=\"ru\">\n<head>\n<meta charset=\"utf-8\">\n<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
CSS = '<link rel="stylesheet" href="../styles.css">\n'
CSS_ROOT = '<link rel="stylesheet" href="styles.css">\n'
FOOT = "</body>\n</html>\n"

def topnav(prefix: str, active: str) -> str:
    def cls(key: str) -> str:
        return ' class="cur"' if active == key else ""
    q = '<a href="%svoprosy-kandidatam.html"%s>Все вопросы</a>' % (prefix, cls("v"))
    d = '<a href="%sindex.html"%s>Все кандидаты</a>' % (prefix, cls("d"))
    return f'<nav class="topnav">{d}{q}</nav>\n'


def make_link(url: str) -> str:
    url = (url or "").strip()
    if not url:
        return ""
    if "://" in url:
        href = url
        disp = re.sub(r"^https?://", "", url)
    else:
        href = "https://" + url
        disp = url
    disp = disp.rstrip("/")
    return f'<a href="{html.escape(href, quote=True)}" target="_blank" rel="noopener">{html.escape(disp)}</a>'


def contacts_span(cells) -> str:
    links = []
    for cell in cells:
        if not cell:
            continue
        for part in str(cell).split(";"):
            link = make_link(part)
            if link:
                links.append(link)
    if not links:
        return ""
    return ' <span class="contacts"> — ' + ", ".join(links) + "</span>"


def read_questions() -> list[str]:
    with open(QUESTIONS_FILE, encoding="utf-8") as f:
        content = f.read()
    return [b.strip() for b in re.split(r"\n-{3,}\n", content) if b.strip()]


def questions_page() -> str:
    title = "Вопросы кандидатам"
    items = []
    for i, q in enumerate(read_questions(), 1):
        items.append(
            '<section class="question">\n'
            f'  <div class="qhead"><span class="qnum">Вопрос {i}</span>'
            f'<button type="button" class="copy">Скопировать вопрос</button></div>\n'
            f'  <pre class="qtext">{html.escape(q)}</pre>\n'
            "</section>"
        )
    body = (
        topnav("", "v")
        + f"<h1>{title}</h1>\n"
        + "\n".join(items)
        + "\n"
        '<script src="questions.js"></script>\n'
    )
    return wrap(title + " — Выборы", CSS_ROOT, body)
def greet_from_name(full_name: str) -> str:
    parts = full_name.split()
    return " ".join(parts[1:]) if len(parts) > 1 else full_name


def region_page(region: str, rows) -> str:
    title = html.escape(region)
    items = []
    for r in rows:
        full = r[col["Кандидат"]] or ""
        name = html.escape(full)
        party = html.escape(r[col["Партия"]] or "")
        contacts = contacts_span([r[col[c]] for c in CONTACT_COLS])
        greet = html.escape(greet_from_name(full), quote=True)
        items.append(
            f'  <li><span class="name">{name}</span> <span class="party">({party})</span>{contacts}'
            f' <button type="button" class="ask" data-greet="{greet}">Скопировать текст Вашего вопроса</button></li>'
        )
    body = (
        topnav("../", "d")
        + f"<h1>{title}</h1>\n"
        + "<ul>\n"
        + "\n".join(items)
        + "\n</ul>\n"
        '<script src="../voprosy.js"></script>\n'
        '<script src="../questions.js"></script>\n'
    )
    return (
        HEAD + f"<title>{title} — Выборы</title>\n" + CSS + "</head>\n<body>\n" + body + FOOT
    )


def read_region_names() -> dict:
    slug = {}
    for path in glob.glob("regions/*.html"):
        name = None
        with open(path, encoding="utf-8") as f:
            for line in f:
                m = re.search(r"<h1>(.*?)</h1>", line)
                if m:
                    name = m.group(1).strip()
                    break
        if name:
            slug[name] = os.path.basename(path)
    return slug


def wrap(title: str, css_link: str, inner: str) -> str:
    return (
        HEAD + f"<title>{html.escape(title)}</title>\n" + css_link
        + "</head>\n<body>\n" + inner + FOOT
    )


def main() -> None:
    slug = read_region_names()

    # Question bank as a browser-visible JS data file (shared by region pages).
    bank = read_questions()
    with open("voprosy.js", "w", encoding="utf-8") as f:
        f.write("window.VOPROSY = " + json.dumps(bank, ensure_ascii=False, indent=1) + ";\n")

    wb = load_workbook(REPO)
    ws = wb[REPO_SHEET]
    headers = [c.value for c in ws[1]]
    global col
    col = {h: i for i, h in enumerate(headers)}

    by_region = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col["Кандидат"]]:
            continue
        if row[col["Партия"]] not in ALLOWED_PARTIES:
            continue
        by_region.setdefault(row[col["Регион"]], []).append(row)

    missing = []
    for region, rows in by_region.items():
        if region not in slug:
            missing.append(region)
            continue
        with open("regions/" + slug[region], "w", encoding="utf-8") as f:
            f.write(region_page(region, rows))

    # «Вопросы кандидатам»: question list with copy buttons.
    with open("voprosy-kandidatam.html", "w", encoding="utf-8") as f:
        f.write(questions_page())

    # Root page: list of all regions.
    nav = []
    for region in sorted(by_region, key=lambda s: s.lower()):
        if region in slug:
            nav.append(f'<li><a href="regions/{slug[region]}">{html.escape(region)}</a></li>')
    root_inner = (
        topnav("", "d")
        + "<h1>Выборы</h1>\n"
        + '<nav class="menu"><ul>\n' + "\n".join(nav) + "\n</ul></nav>\n"
    )
    with open("index.html", "w", encoding="utf-8") as f:
        f.write(wrap("Выборы", CSS_ROOT, root_inner))
    print(f"regions={len(by_region)} nav={len(nav)} questions={len(read_questions())}")
    if missing:
        print("no slug for:", missing)


if __name__ == "__main__":
    main()
